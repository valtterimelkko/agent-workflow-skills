#!/usr/bin/env python3
"""
Validation script for PPTX draft conversions.

This script compares an original .pptx file with a converted _FINAL.pptx file
to verify conversion quality and identify any issues.

Usage:
    python3 pptx_draft_validator.py original.pptx final.pptx [--verbose]
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class PptxValidator:
    """Validate PowerPoint conversion quality."""

    def __init__(self, original_path, final_path, verbose=False):
        self.original_path = Path(original_path)
        self.final_path = Path(final_path)
        self.verbose = verbose
        self.issues = []
        self.warnings = []
        self.info = []

    def validate(self):
        """Run full validation."""
        print(f"🔍 Validating PowerPoint conversion")
        print(f"   Original: {self.original_path.name}")
        print(f"   Final:    {self.final_path.name}")
        print()

        # Load presentations
        try:
            self.original_prs = Presentation(str(self.original_path))
            self.final_prs = Presentation(str(self.final_path))
        except Exception as e:
            print(f"❌ ERROR: Failed to load presentations: {e}")
            return False

        # Run validation checks
        self._check_slide_count()
        self._check_layouts()
        self._check_content_structure()
        self._check_visual_elements()

        # Report results
        self._print_report()

        return len(self.issues) == 0

    def _check_slide_count(self):
        """Check if slide counts match expectations."""
        original_count = len(self.original_prs.slides)
        final_count = len(self.final_prs.slides)

        if final_count < original_count:
            self.issues.append(
                f"Slide count decreased: {original_count} → {final_count}"
            )
        elif final_count > original_count:
            self.info.append(
                f"Slide count increased: {original_count} → {final_count} "
                f"(+{final_count - original_count} new slides)"
            )
        else:
            self.info.append(f"Slide count unchanged: {original_count} slides")

    def _check_layouts(self):
        """Check if layouts are preserved/appropriate."""
        # Check if all original slide layouts are preserved
        for idx, (orig_slide, final_slide) in enumerate(
            zip(self.original_prs.slides, self.final_prs.slides)
        ):
            orig_layout = orig_slide.slide_layout.name
            final_layout = final_slide.slide_layout.name

            if orig_layout != final_layout:
                self.warnings.append(
                    f"Slide {idx}: Layout changed from '{orig_layout}' to '{final_layout}'"
                )

        # Check layouts of new slides (if any)
        original_count = len(self.original_prs.slides)
        final_count = len(self.final_prs.slides)

        if final_count > original_count:
            for idx in range(original_count, final_count):
                layout_name = self.final_prs.slides[idx].slide_layout.name
                self.info.append(f"Slide {idx} (new): Using layout '{layout_name}'")

    def _check_content_structure(self):
        """Check if content structure is preserved."""
        for idx, (orig_slide, final_slide) in enumerate(
            zip(self.original_prs.slides, self.final_prs.slides)
        ):
            # Count text frames
            orig_textboxes = sum(1 for shape in orig_slide.shapes if hasattr(shape, "text_frame"))
            final_textboxes = sum(1 for shape in final_slide.shapes if hasattr(shape, "text_frame"))

            if orig_textboxes != final_textboxes:
                self.warnings.append(
                    f"Slide {idx}: Text box count changed: {orig_textboxes} → {final_textboxes}"
                )

            # Count tables
            orig_tables = sum(1 for shape in orig_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE)
            final_tables = sum(1 for shape in final_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE)

            if orig_tables != final_tables:
                self.warnings.append(
                    f"Slide {idx}: Table count changed: {orig_tables} → {final_tables}"
                )

            # Check for very long text (potential cramming)
            for shape_idx, shape in enumerate(final_slide.shapes):
                if hasattr(shape, "text_frame"):
                    text = shape.text
                    if len(text) > 1000:  # Arbitrary threshold
                        self.warnings.append(
                            f"Slide {idx}, TextBox {shape_idx}: Very long text ({len(text)} chars) - possible cramming"
                        )

    def _check_visual_elements(self):
        """Check if visual elements are preserved."""
        for idx, (orig_slide, final_slide) in enumerate(
            zip(self.original_prs.slides, self.final_prs.slides)
        ):
            # Count images
            orig_images = sum(
                1 for shape in orig_slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            )
            final_images = sum(
                1 for shape in final_slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            )

            if orig_images != final_images:
                self.issues.append(
                    f"Slide {idx}: Image count changed: {orig_images} → {final_images}"
                )

            # Count total shapes as a rough check
            orig_shape_count = len(orig_slide.shapes)
            final_shape_count = len(final_slide.shapes)

            if self.verbose and orig_shape_count != final_shape_count:
                self.info.append(
                    f"Slide {idx}: Shape count changed: {orig_shape_count} → {final_shape_count}"
                )

    def _print_report(self):
        """Print validation report."""
        print("=" * 70)
        print("VALIDATION REPORT")
        print("=" * 70)
        print()

        if self.issues:
            print("❌ ISSUES (must fix):")
            for issue in self.issues:
                print(f"   • {issue}")
            print()

        if self.warnings:
            print("⚠️  WARNINGS (review carefully):")
            for warning in self.warnings:
                print(f"   • {warning}")
            print()

        if self.info:
            print("ℹ️  INFO:")
            for info_item in self.info:
                print(f"   • {info_item}")
            print()

        # Summary
        if not self.issues and not self.warnings:
            print("✅ VALIDATION PASSED - No issues or warnings detected")
        elif not self.issues:
            print(f"⚠️  VALIDATION PASSED WITH WARNINGS - Review {len(self.warnings)} warning(s)")
        else:
            print(f"❌ VALIDATION FAILED - {len(self.issues)} issue(s) detected")

        print()


def main():
    """Main entry point."""
    import argparse

    parser = argparse.ArgumentParser(
        description='Validate PowerPoint draft conversion quality'
    )
    parser.add_argument('original', help='Original .pptx file')
    parser.add_argument('final', help='Final converted .pptx file')
    parser.add_argument('--verbose', '-v', action='store_true',
                       help='Show detailed information')

    args = parser.parse_args()

    validator = PptxValidator(args.original, args.final, verbose=args.verbose)
    success = validator.validate()

    sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()
